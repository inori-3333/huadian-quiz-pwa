#!/usr/bin/env python3
"""Import the 18 week-three worksheets (requires openpyxl for reading only)."""

import argparse
import json
import re
from collections import Counter
from pathlib import Path

import openpyxl

from parse_week2_safety_export import normalize_judge


def clean(value):
    return re.sub(r"\s+", " ", str(value) if value is not None else "").strip()


def convert(source):
    workbook = openpyxl.load_workbook(source, data_only=True)
    names = [f"{n}组" for n in ["一", "二", "三", "四", "五", "六", "七", "八", "九", "十", "十一", "十二", "十三", "十四", "十五", "十六", "十七", "十八"]]
    if workbook.sheetnames != names:
        raise ValueError(f"Unexpected groups: {workbook.sheetnames}")
    questions = []
    for group, sheet in enumerate(workbook, 1):
        rows = list(sheet.values)
        if list(rows[0][:3]) != ["题型", "题干", "答案"] or len(rows) != 16:
            raise ValueError(f"Unexpected layout: {sheet.title}")
        section = None
        counts = Counter()
        for row_number, row in enumerate(rows[1:], 2):
            section = clean(row[0]) or section
            expected = ["选择", "填空", "判断"][(row_number - 2) // 5]
            if section != expected:
                raise ValueError(f"Unexpected type: {sheet.title}:{row_number}")
            stem, raw = clean(row[1]), clean(row[2])
            answer, options = raw, []
            if section == "选择":
                markers = list(re.finditer(r"(?<![A-Z])([A-D])(?:[.．、:：）)]|\s+|(?=[\u4e00-\u9fff]))", stem))
                if [m.group(1) for m in markers] != list("ABCD"):
                    raise ValueError(f"Invalid options: {sheet.title}:{row_number}: {stem}")
                for index, match in enumerate(markers):
                    end = markers[index + 1].start() if index < 3 else len(stem)
                    options.append({"key": match.group(1), "text": stem[match.end():end].strip()})
                stem = stem[:markers[0].start()].strip()
                if answer not in "ABCD" or len(answer) != 1 or any(not o["text"] for o in options):
                    raise ValueError(f"Invalid choice: {sheet.title}:{row_number}")
            elif section == "判断":
                answer = normalize_judge(raw)
            if not stem or not answer:
                raise ValueError(f"Missing question/answer: {sheet.title}:{row_number}")
            counts[section] += 1
            questions.append({
                "id": f"safetyweek3-{len(questions) + 1:04d}",
                "number": row_number - 1,
                "chapter": f"第{group}组", "section": section + "题",
                "type": {"选择": "single", "填空": "fill", "判断": "judge"}[section],
                "stem": stem, "options": options, "answer": answer, "answerRaw": raw,
                "sourceSheet": sheet.title, "sourceRow": row_number,
            })
        if counts != Counter({"选择": 5, "填空": 5, "判断": 5}):
            raise ValueError(f"Unexpected counts: {sheet.title}: {counts}")
    return {"id": "safetyweek3", "title": "第三周安规考试",
            "source": "用户提供的《第三周安规题库汇总.xlsx》，18组题目完整合并，保留原题及答案",
            "questionCount": len(questions), "chapters": [f"第{g}组" for g in range(1, 19)],
            "questions": questions}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    bank = convert(args.source)
    args.output.write_text(json.dumps(bank, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Imported {bank['questionCount']} questions from 18 groups")
